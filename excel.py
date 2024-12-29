import openpyxl

def  main():

    wb = openpyxl.load_workbook(f'excel.xlsm')
    sh = wb['Sheet1']

    sum = 0
    print('data1 = ',sh.cell(row=4, column=2).value)
    sum = int(sh.cell(row=4, column=2).value) + int(sh.cell(row=4, column=3).value)
    sh.cell(row=4, column=3).value = sum
