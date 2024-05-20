#
# 宝ヶ池ゴルフクラブ　ICカード申込書印刷アプリ
# 基準用紙を元に指定された番号を付記しながら指定枚数を出力する
# フォーマット変更 2022/6/29
# フォーマット変更 2022/8/10
#
import openpyxl
from openpyxl.styles import Font
from openpyxl.styles import Alignment
from openpyxl.styles.borders import Border, Side
import datetime
import time

ICCARD_NUMBER = '0005001'   #開始カード番号
DATA_NUMBER = 2000     #出力データ件数


line_thin = Side(style='thin', color='000000')  #細い黒い線
line_medium = Side(style='medium', color='000000')  #太い黒い線
line_dashdot = Side(style='dashDot', color='000000')  #細い一点破線

        
        
# 上左を太線で囲う
border_aro_MLT = Border(top=line_medium, left=line_medium) 
# 上を太線で囲う、左は細線
border_aro_MT_TL = Border(top=line_medium, left=line_thin) 
# 左を太線で囲う、上は細線
border_aro_ML_TT = Border(top=line_thin, left=line_medium) 
# 左上を細線で囲う
border_aro_TLT = Border(top=line_thin, left=line_thin) 
# 左下を太線で囲う
border_aro_MLB = Border(bottom=line_medium, left=line_medium) 
# 上左を一点破線で囲う
border_aro_DLT = Border(top=line_dashdot, left=line_dashdot) 
# 下左を一点破線で囲う
border_aro_DLB = Border(bottom=line_dashdot, left=line_dashdot) 


# 上右を太線で囲う
border_aro_MRT = Border(top=line_medium, right=line_medium) 
# 右を下太線で囲う
border_aro_MRB = Border(bottom=line_medium,right=line_medium) 
# 右を太線で囲う、上は細線
border_aro_MR_TT = Border(top=line_thin, right=line_medium) 
# 上右一点破線で囲う
border_aro_DRT = Border(top=line_dashdot, right=line_dashdot) 
# 下右一点破線で囲う
border_aro_DRB = Border(bottom=line_dashdot, right=line_dashdot) 


    
# 上を太線
border_tb_MT = Border(top=line_medium) 
# 上を細線
border_tb_TT = Border(top=line_thin)
# 上を一点破線
border_tb_DT = Border(top=line_dashdot)
# 下を太線
border_tb_MB = Border(bottom=line_medium)
# 下を細線
border_tb_TB = Border(bottom=line_thin)
# 下を一点破線
border_tb_DB = Border(bottom=line_dashdot) 
# 左を太線
border_tb_ML = Border(left=line_medium)    
# 左を一点破線
border_tb_DL = Border(left=line_dashdot)    
# 右を太線
border_tb_MR = Border(right=line_medium) 
# 右を一点破線
border_tb_DR = Border(right=line_dashdot)                               

def lineset(row,col): 
    
    #1行目 管理番号
    if row == 1:
        if  col == 14:
            return   border_tb_TB   
    
    #3行目 ICカード登録用紙
    if row == 3:
        if  col == 2:
            return   border_aro_MLT
        else:
            if  col == 14:
                return   border_aro_MRT         
            else:
                if col != 1:
                    return border_tb_MT
            
    #4行目 お名前
    if row == 4:
        if  col == 2 or col == 3:
            return border_aro_MLT
        else:
            if  col == 14:
                return border_aro_MRT           
            else:
                if col != 1:
                    return border_tb_MT    
    
    #5行目 郵便番号
    if row == 5:
        if  col == 2 or col == 3:
            return   border_aro_ML_TT
        else:           
            if  col == 14:
                return   border_aro_MR_TT         
            else:
                if col != 1:
                    return border_tb_TT
    
    #6行目 ご住所（1行目）
    if row == 6:
        if  col == 2 or col == 3: 
            return   border_aro_ML_TT
        else:
            if  col == 14:
                return border_aro_MR_TT     
            else:
                if col != 1:
                    return border_tb_TT 
    
    #7行目 ご住所（2行目）
    if row == 7: 
        if  col == 2 or col == 3: 
            return   border_tb_ML
        else:
            if  col == 14:
                return border_tb_MR     
    
    #8行目 お電話番号
    if row == 8: 
        if  col == 2 or col == 3: 
            return border_aro_ML_TT
        else:
            if  col == 14:
                return border_aro_MR_TT 
            else:
                if col != 1:
                    return border_tb_TT
    
    #9行目 生年月日
    if row == 9: 
        if  col == 2 or col == 3: 
            return border_aro_ML_TT
        else:
            if  col == 14:
                return border_aro_MR_TT 
            else:                
                if col != 1:                
                    return border_tb_TT
    
    #10行目 性別
    if row == 10: 
        if  col == 2 or col == 3: 
            return border_aro_ML_TT
        else:
            if  col == 14:
                return border_aro_MR_TT 
            else:
                if col != 1:
                    return border_tb_TT
    
    #11行目 社内情報欄
    if row == 11: 
        if  col == 2:
            return border_aro_MLT
        else:
            if  col == 14:
                return border_aro_MRT
            else:
                if col != 1:
                    return border_tb_MT
            
    #12行目 カード番号
    if row == 12: 
        if  col == 2 or col == 3: 
            return border_aro_MLT
        else:
            if  col == 14:
                return border_aro_MRT
            else:
                if  col >= 4 and col <=12:
                    return border_aro_MT_TL
                else:
                    if col != 1:
                        return border_tb_MT 
                 
    #13行目 登録用紙お渡し日
    if row == 13: 
        if  col == 2 or col == 3:
            return border_aro_ML_TT
        else:
            if  col == 14:
                return border_aro_MR_TT
            else:
                if col != 1:
                    return border_tb_TT   
    
    #14行目 情報入力日
    if row == 14: 
        if  col == 2 or col == 3:
            return border_aro_ML_TT
        else:
            if  col == 14:
                return border_aro_MR_TT
            else:
                if col != 1:
                    return border_tb_TT   
    
    #15行目 カードお渡し日
    if row == 15: 
        if  col == 2 or col == 3:
            return border_aro_ML_TT
        else:
            if  col == 14:
                return border_aro_MR_TT
            else:
                if col != 1:
                    return border_tb_TT   
    
    #16行目 回数券移行残高
    if row == 16: 
        if  col == 2 or col == 3:
            return border_aro_ML_TT
        else:
            if  col == 14:
                return border_aro_MR_TT
            else:
                if col != 1:
                    return border_tb_TT   
    
    #17行目 顧客種別（1行目)
    if row == 17: 
        if  col == 2 or col == 3:
            return border_aro_ML_TT
        else:
            if  col == 14:
                return border_aro_MR_TT
            else:
                if col != 1:
                    return border_tb_TT   
    
    #18行目 顧客種別（2行目)
    if row == 18: 
        if  col == 2 or col == 3:
            return   border_tb_ML
        else:
            if  col == 14:
                return border_tb_MR
    
    #19行目 ロッカー契約（1行目)
    if row == 19: 
        if  col == 2 or col == 3:
            return border_aro_ML_TT
        else:
            if  col == 14:
                return border_aro_MR_TT
            else:
                if col != 1:
                    return border_tb_TT   
    
    #20行目 ロッカー契約（2行目)
    if row == 20: 
        if  col == 2 or col == 3:
            return   border_tb_ML
        else:
            if  col == 14:
                return border_tb_MR
    
    #21行目 摘要（1行目)
    if row == 21: 
        if  col == 2 or col == 3:
            return border_aro_ML_TT
        else:
            if  col == 14:
                return border_aro_MR_TT
            else:
                if col != 1:
                    return border_tb_TT   
    
    #22行目 摘要（2行目)
    if row == 22: 
        if  col == 2 or col == 3:
            return   border_aro_MLB
        else:
            if  col == 14:
                return border_aro_MRB
            else:
                if col != 1:
                    return border_tb_MB
    
    #30行目 貼り付け欄
    if row == 30:
        if col == 9:
            return  border_aro_DLT
        if col == 14:
            return  border_aro_DRT
        if col  >=10 and col <=13:
            return  border_tb_DT
        
    if row >= 31 and row <= 39:
        if col == 9:
            return  border_tb_DL
        if col == 14:
            return  border_tb_DR
    
    if row == 40:
        if col == 9:
            return  border_aro_DLB
        if col == 14:
            return  border_aro_DRB
        if col  >=10 and col <=13:
            return  border_tb_DB    
    
    #57行目　署名欄             
    if row == 57: 
        if col >=5 and col <= 15:
            return border_tb_MB    

# 次のカード番号を配列で返す
def number_calc(cardno_now):    
    iccard_next_num = ['0','0','0','0','0','0','0']
    card_num_len = len(str(cardno_now))
    
    # カードナンバーを配列にセット
    i = 6
    idx =  card_num_len - 1
    while i <= 6 and idx >= 0:
        iccard_next_num[i] = str(cardno_now)[idx:idx+1]
        i -= 1
        idx -= 1    
        
    return iccard_next_num              

if __name__ == "__main__":
    wb = openpyxl.load_workbook('C:\entry\entrysheet.xlsx')
    ws = wb['Sheet0']
    
    dt_now = datetime.datetime.now()
    ut = time.time()
    print('処理開始：',dt_now)    

    data_cnt_now = 1    
    cardno_now = int(ICCARD_NUMBER)
    
    while data_cnt_now <=  DATA_NUMBER:        
        # 1行目から順番にコピーする
        for data_row in range(1,58):
            #コピー先行番号取得
            copyrow = data_row + data_cnt_now*57
            for col in range(1,15):
                ws.row_dimensions[copyrow].height =  ws.row_dimensions[data_row].height #行の高さ合わせ                                 
                ws.cell(row=copyrow, column=col).value =  ws.cell(row=data_row, column=col).value #セル内容のコピー
                # カード番号のセット
                if data_row == 12:
                    if col == 5:
                        card_no = []
                        card_no = number_calc(cardno_now)
                        cardno_now += 1
                        idx = 0
                    if col == 3:
                        ws.cell(row=copyrow, column=col).value = 'E'
                    if col == 4:
                        ws.cell(row=copyrow, column=col).value = 'S'
                    if col >= 5 and col <= 11:
                        ws.cell(row=copyrow, column=col).value = card_no[idx]
                        idx += 1
                                                    
                #罫線を引く            
                res_border = lineset(data_row,col)
                ws.cell(row=copyrow, column=col).border = res_border
                
                font = ws.cell(row=data_row, column=col).font #フォントを合わせる
                ws.cell(row=copyrow, column=col).font = Font(name = font.name, color = '000000', bold=font.bold,size = font.sz)
                alignment = ws.cell(row=data_row, column=col).alignment #文字位置を合わせる
                ws.cell(row=copyrow, column=col).alignment = Alignment(horizontal = alignment.horizontal, vertical = alignment.vertical)
        
        data_cnt_now += 1          
           
    #保存
    wb.save('C:\entry\entrysheet.xlsx')
    
    dt_now = datetime.datetime.now()
    t = time.time() - ut
    td = datetime.timedelta(seconds=t)
    print('処理終了：',dt_now)
    print('入力件数：',data_cnt_now-1)
    print(td) 

   