#################################################
#
#   営業提出資料
#   納品書データ取得
#  
#
#   指定期間に作成された納品書を取得し売上実績を把握
#
#   2024_1_24 新規作成
#   python 3.10.3
#
#################################################

import pyodbc as pyo
import openpyxl
from datetime import date

#ワークシートを事前にクリアする
wb = openpyxl.load_workbook(r'C:\Users\hishi\OneDrive\Workplace\2024年営業計画\wk_nouhin.xlsx')
wb.remove(wb.worksheets[-1])
ws = wb.create_sheet(title="Sheet1")
wb.save(r'C:\Users\hishi\OneDrive\Workplace\2024年営業計画\wk_nouhin.xlsx')

##########################################
#
#  納品書DB　アクセス
#
##########################################

con_str1 = (
	r'Driver={Microsoft Access Driver (*.mdb, *.accdb)};'
	r'DBQ=C:\Users\hishi\OneDrive\Workplace\2024年営業計画\2024-07月納品データ\納品書Ver.1.2元データ用(2024版 税率10% 西暦表示).mdb;'
	)

con = pyo.connect(con_str1)
cursor = con.cursor()
###################################
# 納品書データを日付で抽出
# 指定日付はSQL文を変更
# ###################################
sql1 = 'SELECT * FROM 納品書 \
        LEFT JOIN ゴルフ場名簿 ON(納品書.ゴルフ場No = ゴルフ場名簿.ゴルフ場No) \
        WHERE 納品日 Between #2024/07/01# AND #2024/07/31#'
#sql1 = 'SELECT * FROM 納品書 \
#        LEFT JOIN ゴルフ場名簿 ON(納品書.ゴルフ場No = ゴルフ場名簿.ゴルフ場No) \
#        WHERE 納品No > 6589'
rows_nouhin = cursor.execute(sql1).fetchall()

#print(f'納品No={row.納品No}, 納品合計金額={row.納品合計金額}, 摘要={row.摘要}, 納品日={row.納品日.strftime("%Y/%m/%d")}')
print('納品書抽出処理開始')

rows_nouhin_len = len(rows_nouhin)
###################################
#  売上実績をワークファイルに書き込み
###################################
#wb = openpyxl.load_workbook(r'C:\Users\user\OneDrive\Workplace\2024年営業計画\売上計画案（東京本社）.xlsx')
wb = openpyxl.load_workbook(r'C:\Users\hishi\OneDrive\Workplace\2024年営業計画\wk_nouhin.xlsx')
sh_nouhin = wb['Sheet1']
rowno = sh_nouhin.max_row + 1
start_rowno = rowno
#print(f'最終行 = {maxr}')
data_no = 0
for data_no in range(0, rows_nouhin_len):
    if rows_nouhin[data_no].納品合計金額 != None:
        if rows_nouhin[data_no].納品合計金額 != 0:
            sh_nouhin.cell(rowno,1).value = rows_nouhin[data_no].納品No
            #sh_nouhin.cell(rowno,2).value = rows_nouhin[data_no].納品日
            w_date = date(int(rows_nouhin[data_no].納品日.year), int(rows_nouhin[data_no].納品日.month), int(rows_nouhin[data_no].納品日.day))
            sh_nouhin.cell(rowno,2).value = w_date
            sh_nouhin.cell(rowno,3).value = rows_nouhin[data_no].ゴルフ場名
            sh_nouhin.cell(rowno,4).value = rows_nouhin[data_no].摘要
            sh_nouhin.cell(rowno,5).value = rows_nouhin[data_no].納品合計金額 
            sh_nouhin.cell(rowno,6).value = w_date.year
            sh_nouhin.cell(rowno,7).value = w_date.month
            rowno += 1
            
#wb.save(r'C:\Users\user\OneDrive\Workplace\2024年営業計画\売上計画案（東京本社）.xlsx')
wb.save(r'C:\Users\hishi\OneDrive\Workplace\2024年営業計画\wk_nouhin.xlsx')

cursor.close()
con.close()

